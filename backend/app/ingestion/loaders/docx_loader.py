"""
DOCX document loader.

Reads Microsoft Word (.docx) files and converts them
into the standardized IngestedDocument format.
"""

from pathlib import Path

from docx import Document

from app.ingestion.loaders.base_loader import BaseLoader
from app.ingestion.schemas.document_schema import IngestedDocument


class DOCXLoader(BaseLoader):
    """
    Loader for Microsoft Word documents.
    """

    def load(self, file_path: Path) -> IngestedDocument:
        """
        Load a DOCX document.
        """

        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        document = Document(file_path)

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        content = "\n".join(paragraphs)

        return IngestedDocument(
            source="docx",
            source_file=file_path.name,
            content=content,
            metadata={
                "extension": file_path.suffix,
                "paragraphs": len(document.paragraphs),
                "file_size": file_path.stat().st_size,
            },
        )